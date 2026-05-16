"""
文档加载器测试脚本。

功能：
1. 自动生成测试用 PDF / DOCX / MD 样本文件
2. 逐一验证 DocumentPipeline 能否正确解析
3. 验证分块结果、元数据完整性
4. 验证异常处理（不支持格式、文件不存在）
"""

import logging
import sys
from pathlib import Path

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s - %(message)s",
    stream=sys.stdout,
)
logger = logging.getLogger("test")

TEST_DIR = Path(__file__).parent / "test_samples"


# ============================================================
# 1. 生成测试样本文件
# ============================================================

def create_test_markdown() -> Path:
    """生成测试用 Markdown 文件"""
    path = TEST_DIR / "test_sample.md"
    path.write_text(
        "# 企业知识库助手\n\n"
        "## 项目简介\n"
        "本项目基于 RAG 架构，实现企业内部知识库的智能问答。\n\n"
        "## 技术栈\n"
        "- FastAPI\n"
        "- LangChain\n"
        "- Chroma\n\n"
        "## 使用说明\n"
        "请将文档上传至系统，系统会自动解析并建立索引。\n"
        "用户可以通过 API 接口进行问答查询。\n",
        encoding="utf-8",
    )
    return path


def create_test_docx() -> Path:
    """生成测试用 DOCX 文件"""
    from docx import Document

    path = TEST_DIR / "test_sample.docx"
    doc = Document()
    doc.add_heading("员工手册", level=1)
    doc.add_paragraph("第一章 公司简介")
    doc.add_paragraph("本公司成立于2020年，专注于人工智能技术的研发与应用。")
    doc.add_paragraph("第二章 考勤制度")
    doc.add_paragraph("工作时间为每周一至周五，上午9:00至下午18:00。")
    doc.add_paragraph("迟到超过30分钟视为旷工半天。")
    doc.add_heading("第三章 假期制度", level=1)
    doc.add_paragraph("年假根据工龄计算，满1年可享受5天年假。")
    doc.add_paragraph("病假需提供医院证明，每年累计不超过15天。")
    doc.save(str(path))
    return path


def create_test_pdf() -> Path:
    """生成测试用 PDF 文件"""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = TEST_DIR / "test_sample.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4

    c.setFont("Helvetica-Bold", 18)
    c.drawString(100, height - 100, "Product Manual")

    c.setFont("Helvetica", 12)
    lines = [
        "Chapter 1: Introduction",
        "This product is an enterprise knowledge base assistant.",
        "It uses RAG (Retrieval-Augmented Generation) technology.",
        "",
        "Chapter 2: Features",
        "- Document upload and parsing (PDF, DOCX, Markdown)",
        "- Semantic search powered by Chroma vector database",
        "- Multi-model support (DeepSeek, Xiaomi MiMo, local models)",
        "",
        "Chapter 3: Getting Started",
        "1. Upload your documents via the API",
        "2. Wait for the system to process and index them",
        "3. Start asking questions through the chat interface",
    ]
    for i, line in enumerate(lines):
        c.drawString(100, height - 140 - i * 24, line)

    c.save()
    return path


def setup_test_files() -> dict[str, Path]:
    """生成所有测试样本，返回 {格式: 路径} 映射"""
    TEST_DIR.mkdir(parents=True, exist_ok=True)

    files = {}
    files["md"] = create_test_markdown()
    logger.info("已生成 Markdown 测试文件: %s", files["md"])

    files["docx"] = create_test_docx()
    logger.info("已生成 DOCX 测试文件: %s", files["docx"])

    files["pdf"] = create_test_pdf()
    logger.info("已生成 PDF 测试文件: %s", files["pdf"])

    return files


# ============================================================
# 2. 测试用例
# ============================================================

REQUIRED_METADATA_KEYS = {"source", "filename", "file_type", "created_at", "source_path"}


def _assert_metadata(chunks: list, file_path: Path, expected_type: str):
    """验证所有 chunk 的 metadata 完整性"""
    for i, chunk in enumerate(chunks):
        meta = chunk.metadata
        assert chunk.page_content, f"第 {i} 个 chunk 内容为空"

        missing = REQUIRED_METADATA_KEYS - set(meta.keys())
        assert not missing, f"第 {i} 个 chunk 缺少 metadata 字段: {missing}"

        assert meta["source"] == file_path.name, (
            f"第 {i} 个 chunk source 不正确: {meta['source']}"
        )
        assert meta["filename"] == file_path.name
        assert meta["file_type"] == expected_type
        assert meta["source_path"] == str(file_path.resolve())
        assert meta["created_at"], "created_at 为空"


def test_markdown(pipeline, file_path: Path):
    """测试 Markdown 解析"""
    logger.info("=" * 50)
    logger.info("测试 Markdown 解析: %s", file_path.name)

    chunks = pipeline.process(file_path)
    assert len(chunks) > 0, "Markdown 解析结果为空"
    _assert_metadata(chunks, file_path, ".md")

    logger.info("Markdown 测试通过: %d 个 chunk", len(chunks))
    _print_chunks_preview(chunks)


def test_docx(pipeline, file_path: Path):
    """测试 DOCX 解析"""
    logger.info("=" * 50)
    logger.info("测试 DOCX 解析: %s", file_path.name)

    chunks = pipeline.process(file_path)
    assert len(chunks) > 0, "DOCX 解析结果为空"
    _assert_metadata(chunks, file_path, ".docx")

    logger.info("DOCX 测试通过: %d 个 chunk", len(chunks))
    _print_chunks_preview(chunks)


def test_pdf(pipeline, file_path: Path):
    """测试 PDF 解析"""
    logger.info("=" * 50)
    logger.info("测试 PDF 解析: %s", file_path.name)

    chunks = pipeline.process(file_path)
    assert len(chunks) > 0, "PDF 解析结果为空"
    _assert_metadata(chunks, file_path, ".pdf")

    logger.info("PDF 测试通过: %d 个 chunk", len(chunks))
    _print_chunks_preview(chunks)


def test_unsupported_format(pipeline):
    """测试不支持的文件格式"""
    logger.info("=" * 50)
    logger.info("测试不支持的文件格式: .xyz")

    fake_file = TEST_DIR / "test.xyz"
    fake_file.write_text("hello")

    try:
        pipeline.process(fake_file)
        assert False, "应该抛出 ValueError"
    except ValueError as e:
        logger.info("异常处理正确: %s", e)


def test_file_not_found(pipeline):
    """测试文件不存在"""
    logger.info("=" * 50)
    logger.info("测试文件不存在")

    try:
        pipeline.process(TEST_DIR / "nonexistent.pdf")
        assert False, "应该抛出 FileNotFoundError"
    except FileNotFoundError as e:
        logger.info("异常处理正确: %s", e)


def test_batch_process(pipeline, files: dict[str, Path]):
    """测试批量处理"""
    logger.info("=" * 50)
    logger.info("测试批量处理")

    chunks = pipeline.process_batch(list(files.values()))
    assert len(chunks) > 0, "批量处理结果为空"

    sources = {c.metadata.get("source") for c in chunks}
    assert len(sources) == len(files), f"期望 {len(files)} 个文件的 chunk，实际来源: {sources}"

    logger.info("批量处理测试通过: %d 个 chunk，来源文件 %s", len(chunks), sources)


def test_splitter_strategies(files: dict[str, Path]):
    """测试不同分块策略"""
    from app.core.document import DocumentPipeline

    # 用较长的文本测试，确保能切出多个 chunk
    long_text_path = TEST_DIR / "long_text.md"
    long_text_path.write_text(
        "\n\n".join([f"## 第{i}章\n" + "这是测试内容。" * 50 for i in range(1, 6)]),
        encoding="utf-8",
    )

    for strategy in ["fixed", "recursive", "smart"]:
        logger.info("=" * 50)
        logger.info("测试分块策略: %s", strategy)

        pipeline = DocumentPipeline(
            chunk_size=200, chunk_overlap=50, splitter_strategy=strategy
        )
        chunks = pipeline.process(long_text_path)
        assert len(chunks) > 1, f"策略 {strategy} 应产生多个 chunk，实际 {len(chunks)}"

        # 验证 metadata 继承
        for chunk in chunks:
            assert chunk.metadata.get("source") == long_text_path.name

        logger.info("策略 %s 测试通过: %d 个 chunk", strategy, len(chunks))


def test_splitter_factory():
    """测试分块器工厂"""
    from app.core.document.splitter import DocumentSplitterFactory

    logger.info("=" * 50)
    logger.info("测试分块器工厂")

    available = DocumentSplitterFactory.available_strategies()
    assert "fixed" in available
    assert "recursive" in available
    assert "smart" in available
    logger.info("可用策略: %s", available)

    # 测试创建各策略
    for strategy in ["fixed", "recursive", "smart"]:
        splitter = DocumentSplitterFactory.create(strategy, chunk_size=200, chunk_overlap=50)
        assert splitter.chunk_size == 200
        assert splitter.chunk_overlap == 50
        logger.info("工厂创建 %s 成功", strategy)

    # 测试不支持的策略
    try:
        DocumentSplitterFactory.create("nonexistent")
        assert False, "应该抛出 ValueError"
    except ValueError as e:
        logger.info("工厂异常处理正确: %s", e)


def test_splitter_invalid_params():
    """测试分块器参数校验"""
    from app.core.document.splitter import DocumentSplitterFactory

    logger.info("=" * 50)
    logger.info("测试分块器参数校验")

    # chunk_size <= 0
    try:
        DocumentSplitterFactory.create("fixed", chunk_size=0, chunk_overlap=10)
        assert False, "应该抛出 ValueError"
    except ValueError:
        logger.info("chunk_size=0 校验正确")

    # overlap >= size
    try:
        DocumentSplitterFactory.create("recursive", chunk_size=100, chunk_overlap=100)
        assert False, "应该抛出 ValueError"
    except ValueError:
        logger.info("overlap >= size 校验正确")

    # overlap < 0
    try:
        DocumentSplitterFactory.create("fixed", chunk_size=100, chunk_overlap=-1)
        assert False, "应该抛出 ValueError"
    except ValueError:
        logger.info("overlap < 0 校验正确")


def test_smart_splitter_priority():
    """测试智能分块器的标题/段落/句子优先级"""
    from langchain_core.documents import Document
    from app.core.document.splitter.smart import SmartSplitter

    logger.info("=" * 50)
    logger.info("测试智能分块器优先级")

    # 构造含标题、段落、句子的长文本
    text = (
        "# 第一章 总则\n\n"
        "第一条 本办法适用于公司全体员工。员工应当遵守公司各项规章制度。\n\n"
        "第二条 公司实行每周五天工作制，每日工作时间为上午九点至下午六点。\n\n"
        "## 第二章 考勤管理\n\n"
        "第三条 员工应当按时上下班，不得迟到早退。迟到超过三十分钟视为旷工半天。\n\n"
        "第四条 员工请假应当提前申请。病假需提供医院证明，事假需经主管审批。\n\n"
        "## 第三章 奖惩制度\n\n"
        "第五条 对工作表现优秀的员工给予表彰和奖励。奖励包括但不限于奖金、晋升和荣誉称号。\n\n"
        "第六条 对违反公司规定的员工视情节轻重给予处分。处分包括警告、记过和解除劳动合同。"
    )

    splitter = SmartSplitter(chunk_size=150, chunk_overlap=20)
    doc = Document(page_content=text, metadata={"source": "test.md"})
    chunks = splitter.split([doc])

    assert len(chunks) > 1, f"智能分块应产生多个 chunk，实际 {len(chunks)}"
    for chunk in chunks:
        assert len(chunk.page_content) <= 150, (
            f"chunk 超过 chunk_size: {len(chunk.page_content)} > 150"
        )
        assert chunk.metadata["source"] == "test.md"

    # 验证按标题切分：章节标题不应被截断
    heading_chunks = [c for c in chunks if "## 第" in c.page_content or "# 第" in c.page_content]
    assert len(heading_chunks) > 0, "应存在以标题开头的 chunk"

    logger.info("智能分块测试通过: %d 个 chunk", len(chunks))
    for i, chunk in enumerate(chunks[:5]):
        preview = chunk.page_content[:60].replace("\n", " ").strip()
        logger.info("  chunk[%d](%d字): %s", i, len(chunk.page_content), preview)


# ============================================================
# 3. 工具函数
# ============================================================

def _print_chunks_preview(chunks, max_preview: int = 3):
    """打印 chunk 预览"""
    for i, chunk in enumerate(chunks[:max_preview]):
        preview = chunk.page_content[:80].replace("\n", " ")
        logger.info("  chunk[%d]: %s... | meta: %s", i, preview, chunk.metadata)
    if len(chunks) > max_preview:
        logger.info("  ... 共 %d 个 chunk", len(chunks))


def cleanup():
    """清理测试生成的文件"""
    import shutil
    if TEST_DIR.exists():
        shutil.rmtree(TEST_DIR)
        logger.info("已清理测试目录: %s", TEST_DIR)


# ============================================================
# 4. 主流程
# ============================================================

def main():
    from app.core.document import DocumentPipeline

    pipeline = DocumentPipeline(chunk_size=200, chunk_overlap=20)

    logger.info("支持的文件格式: %s", pipeline.supported_extensions())
    logger.info("可用分块策略: %s", pipeline.available_strategies())

    try:
        # 生成测试文件
        files = setup_test_files()

        # 逐格式测试
        test_markdown(pipeline, files["md"])
        test_docx(pipeline, files["docx"])
        test_pdf(pipeline, files["pdf"])

        # 异常测试
        test_unsupported_format(pipeline)
        test_file_not_found(pipeline)

        # 批量测试
        test_batch_process(pipeline, files)

        # 分块策略测试
        test_splitter_factory()
        test_splitter_invalid_params()
        test_splitter_strategies(files)
        test_smart_splitter_priority()

        logger.info("=" * 50)
        logger.info("全部测试通过！")

    finally:
        cleanup()


if __name__ == "__main__":
    main()
